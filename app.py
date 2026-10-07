# -*- coding: utf-8 -*-
"""实验课签到系统 —— FastAPI 后端（课程为中心的权限模型）。"""
import io
import os
import secrets
import socket
import sqlite3
from datetime import datetime, timedelta
from urllib.parse import quote

from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from starlette.middleware.sessions import SessionMiddleware

import config
import db
import importers

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")

app = FastAPI(title="实验课签到系统")
app.add_middleware(SessionMiddleware, secret_key=config.SECRET_KEY, same_site="lax", max_age=None)


@app.middleware("http")
async def _no_cache_static(request: Request, call_next):
    """页面与静态资源禁用浏览器缓存，改版后普通刷新即可拿到新文件。"""
    response = await call_next(request)
    if request.url.path == "/" or request.url.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return response


@app.on_event("startup")
def _on_startup():
    db.init_db()
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
    except Exception:
        ip = "127.0.0.1"
    print()
    print("=" * 52)
    print("  实验课签到系统已启动")
    print(f"  访问地址：http://{ip}:8000")
    print("  把这个网址发给其他助教（需处于同一 WiFi/内网）")
    print("=" * 52)
    print()


def require_user(request: Request) -> int:
    """登录校验，返回当前登录用户的 id。未登录（或账号已被删除）抛 401。"""
    uid = request.session.get("uid")
    if not uid:
        raise HTTPException(status_code=401, detail="未登录")
    with db.get_db() as conn:
        if not conn.execute("SELECT id FROM users WHERE id = ?", (uid,)).fetchone():
            request.session.clear()
            raise HTTPException(status_code=401, detail="账号不存在，请重新登录")
    return uid


def require_super_admin(request: Request) -> int:
    """全局超管校验（实时读 DB，不信任 session）。非超管抛 403。"""
    uid = require_user(request)
    with db.get_db() as conn:
        row = conn.execute("SELECT role FROM users WHERE id = ?", (uid,)).fetchone()
        if row and row["role"] == "super_admin":
            return uid
    raise HTTPException(status_code=403, detail="需要全局管理员权限")


def require_course_member(request: Request, course_id: int) -> str:
    """课程成员校验（管理员或助教），返回课程内角色。课程不存在 404，非成员 403。"""
    uid = require_user(request)
    with db.get_db() as conn:
        if not conn.execute("SELECT id FROM courses WHERE id = ?", (course_id,)).fetchone():
            raise HTTPException(status_code=404, detail="课程不存在")
        row = conn.execute(
            "SELECT role FROM course_members WHERE course_id = ? AND user_id = ?",
            (course_id, uid),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=403, detail="你不是该课程的成员")
        return row["role"]


def require_course_admin(request: Request, course_id: int) -> str:
    """课程管理员校验。助教/非成员抛 403。"""
    role = require_course_member(request, course_id)
    if role != "admin":
        raise HTTPException(status_code=403, detail="需要课程管理员权限")
    return role


def _now():
    return datetime.now().astimezone().isoformat(timespec="seconds")


# ---------- 登录 / 注册 ----------

@app.post("/api/login")
async def login(request: Request):
    data = await request.json()
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""
    operator = (data.get("operator") or "").strip()
    if not username or not password:
        raise HTTPException(status_code=401, detail="请输入用户名和密码")
    with db.get_db() as conn:
        user = conn.execute(
            "SELECT id, username, display_name, role, password_hash "
            "FROM users WHERE username = ?",
            (username,),
        ).fetchone()
    if not user or not db.verify_password(password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    request.session["uid"] = user["id"]
    request.session["username"] = user["username"]
    request.session["display_name"] = user["display_name"] or user["username"]
    request.session["operator"] = operator or user["display_name"] or user["username"]
    request.session["role"] = user["role"]
    return {
        "ok": True,
        "id": user["id"],
        "username": user["username"],
        "display_name": request.session["display_name"],
        "operator": request.session["operator"],
        "role": user["role"],
    }


@app.post("/api/register")
async def register(request: Request):
    data = await request.json()
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""
    display_name = (data.get("display_name") or "").strip()
    if not username or not password:
        raise HTTPException(status_code=400, detail="用户名和密码不能为空")
    if len(password) < 6:
        raise HTTPException(status_code=400, detail="密码至少 6 位")
    try:
        with db.get_db() as conn:
            cur = conn.execute(
                "INSERT INTO users(username, password_hash, display_name, role, created_at) "
                "VALUES (?, ?, ?, 'user', ?)",
                (username, db.hash_password(password), display_name or None, _now()),
            )
    except sqlite3.IntegrityError:
        raise HTTPException(status_code=400, detail="该用户名已被注册")
    request.session["uid"] = cur.lastrowid
    request.session["username"] = username
    request.session["display_name"] = display_name or username
    request.session["operator"] = display_name or username
    request.session["role"] = "user"
    return {
        "ok": True,
        "id": cur.lastrowid,
        "username": username,
        "display_name": request.session["display_name"],
        "operator": request.session["operator"],
        "role": "user",
    }


@app.post("/api/logout")
async def logout(request: Request):
    request.session.clear()
    return {"ok": True}


@app.get("/api/me")
async def me(request: Request):
    role = "user"
    uid = request.session.get("uid")
    if uid:
        with db.get_db() as conn:
            row = conn.execute("SELECT role FROM users WHERE id = ?", (uid,)).fetchone()
        if row:
            role = row["role"]
    return {
        "authed": bool(request.session.get("uid")),
        "id": uid,
        "username": request.session.get("username", ""),
        "display_name": request.session.get("display_name", ""),
        "operator": request.session.get("operator", ""),
        "role": role,
        "is_super_admin": role == "super_admin",
    }


# ---------- 课程 ----------

@app.get("/api/courses")
async def list_courses(uid: int = Depends(require_user)):
    with db.get_db() as conn:
        rows = conn.execute(
            "SELECT c.id, c.name, c.created_at, m.role "
            "FROM course_members m JOIN courses c ON c.id = m.course_id "
            "WHERE m.user_id = ? ORDER BY c.id",
            (uid,),
        ).fetchall()
    return {"courses": [dict(r) for r in rows]}


@app.post("/api/courses")
async def create_course(payload: dict, uid: int = Depends(require_user)):
    name = (payload.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="课程名不能为空")
    with db.get_db() as conn:
        try:
            cur = conn.execute(
                "INSERT INTO courses(name, owner_id, created_at) VALUES (?, ?, ?)",
                (name, uid, _now()),
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=400, detail="你已创建过同名课程")
        conn.execute(
            "INSERT INTO course_members(course_id, user_id, role, joined_at) VALUES (?, ?, 'admin', ?)",
            (cur.lastrowid, uid, _now()),
        )
    return {"id": cur.lastrowid, "name": name}


@app.get("/api/courses/{course_id}")
async def get_course(course_id: int, role: str = Depends(require_course_member)):
    with db.get_db() as conn:
        row = conn.execute(
            "SELECT id, name FROM courses WHERE id = ?", (course_id,)
        ).fetchone()
        count = conn.execute(
            "SELECT COUNT(*) FROM course_members WHERE course_id = ?", (course_id,)
        ).fetchone()[0]
    return {"id": row["id"], "name": row["name"], "role": role, "member_count": count}


@app.delete("/api/courses/{course_id}")
async def delete_course(course_id: int, role: str = Depends(require_course_admin)):
    with db.get_db() as conn:
        conn.execute("DELETE FROM courses WHERE id = ?", (course_id,))
    return {"ok": True}


# ---------- 成员与邀请 ----------

@app.get("/api/courses/{course_id}/members")
async def list_members(course_id: int, role: str = Depends(require_course_admin)):
    with db.get_db() as conn:
        rows = conn.execute(
            "SELECT m.id, m.user_id, u.username, u.display_name, m.role, m.joined_at "
            "FROM course_members m JOIN users u ON u.id = m.user_id "
            "WHERE m.course_id = ? ORDER BY m.joined_at",
            (course_id,),
        ).fetchall()
    return {"members": [dict(r) for r in rows]}


@app.delete("/api/courses/{course_id}/members/{user_id}")
async def remove_member(course_id: int, user_id: int, role: str = Depends(require_course_admin)):
    with db.get_db() as conn:
        target = conn.execute(
            "SELECT role FROM course_members WHERE course_id = ? AND user_id = ?",
            (course_id, user_id),
        ).fetchone()
        if not target:
            raise HTTPException(status_code=404, detail="该成员不存在")
        if target["role"] == "admin":
            raise HTTPException(status_code=400, detail="不能移除课程创建者")
        conn.execute(
            "DELETE FROM course_members WHERE course_id = ? AND user_id = ?",
            (course_id, user_id),
        )
    return {"ok": True}


def _new_invite_code(conn, course_id, uid, hours):
    """生成一次性邀请码（8 位大写十六进制，无易混字符）。返回 (code, expires_at)。"""
    now = datetime.now().astimezone()
    expires = (now + timedelta(hours=hours)).isoformat(timespec="seconds")
    for _ in range(3):
        code = secrets.token_hex(4).upper()
        try:
            conn.execute(
                "INSERT INTO invitations(course_id, code, created_by, created_at, expires_at) "
                "VALUES (?, ?, ?, ?, ?)",
                (course_id, code, uid, now.isoformat(timespec="seconds"), expires),
            )
            return code, expires
        except sqlite3.IntegrityError:
            continue
    raise HTTPException(status_code=500, detail="邀请码生成失败，请重试")


@app.get("/api/courses/{course_id}/invites")
async def list_invites(course_id: int, role: str = Depends(require_course_admin)):
    with db.get_db() as conn:
        rows = conn.execute(
            "SELECT i.id, i.code, i.created_at, i.expires_at, i.used_at, i.used_by, "
            "(SELECT username FROM users WHERE id = i.used_by) AS used_username "
            "FROM invitations i WHERE i.course_id = ? ORDER BY i.id DESC",
            (course_id,),
        ).fetchall()
    return {"invites": [dict(r) for r in rows]}


@app.post("/api/courses/{course_id}/invites")
async def create_invite(
    payload: dict, request: Request, course_id: int, role: str = Depends(require_course_admin)
):
    hours = payload.get("expires_hours", 168) or 168
    try:
        hours = int(hours)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="有效期格式不正确")
    if not 1 <= hours <= 24 * 365:
        raise HTTPException(status_code=400, detail="有效期需在 1 小时到 365 天之间")
    with db.get_db() as conn:
        code, expires = _new_invite_code(conn, course_id, request.session["uid"], hours)
    return {"code": code, "expires_at": expires}


@app.delete("/api/courses/{course_id}/invites/{invite_id}")
async def revoke_invite(course_id: int, invite_id: int, role: str = Depends(require_course_admin)):
    with db.get_db() as conn:
        cur = conn.execute(
            "DELETE FROM invitations WHERE id = ? AND course_id = ?", (invite_id, course_id)
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="邀请码不存在")
    return {"ok": True}


def _get_invite(conn, code):
    return conn.execute(
        "SELECT i.*, c.name AS course_name, u.username AS admin_username "
        "FROM invitations i JOIN courses c ON c.id = i.course_id "
        "JOIN users u ON u.id = c.owner_id "
        "WHERE i.code = ?",
        (code,),
    ).fetchone()


@app.get("/api/invites/{code}")
async def invite_preview(code: str, uid: int = Depends(require_user)):
    code = (code or "").strip().upper()
    if not code:
        raise HTTPException(status_code=400, detail="请输入邀请码")
    with db.get_db() as conn:
        inv = _get_invite(conn, code)
        if not inv or inv["used_at"] or inv["expires_at"] < _now():
            raise HTTPException(status_code=404, detail="邀请码无效或已过期")
        if conn.execute(
            "SELECT 1 FROM course_members WHERE course_id = ? AND user_id = ?",
            (inv["course_id"], uid),
        ).fetchone():
            return {"already_member": True, "course_name": inv["course_name"]}
    return {
        "already_member": False,
        "course_id": inv["course_id"],
        "course_name": inv["course_name"],
        "role": "assistant",
        "admin_username": inv["admin_username"],
    }


@app.post("/api/invites/accept")
async def accept_invite(payload: dict, uid: int = Depends(require_user)):
    code = (payload.get("code") or "").strip().upper()
    if not code:
        raise HTTPException(status_code=400, detail="请输入邀请码")
    with db.get_db() as conn:
        inv = _get_invite(conn, code)
        if not inv or inv["used_at"] or inv["expires_at"] < _now():
            raise HTTPException(status_code=404, detail="邀请码无效或已过期")
        if conn.execute(
            "SELECT 1 FROM course_members WHERE course_id = ? AND user_id = ?",
            (inv["course_id"], uid),
        ).fetchone():
            raise HTTPException(status_code=400, detail="你已是该课程成员")
        conn.execute(
            "INSERT INTO course_members(course_id, user_id, role, joined_at) "
            "VALUES (?, ?, 'assistant', ?)",
            (inv["course_id"], uid, _now()),
        )
        # 条件 UPDATE 互斥：并发下后到者 rowcount=0，整个事务回滚（含刚插入的成员行）
        cur = conn.execute(
            "UPDATE invitations SET used_at = ?, used_by = ? WHERE id = ? AND used_at IS NULL",
            (_now(), uid, inv["id"]),
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=409, detail="邀请码刚被使用")
    return {"ok": True, "course_id": inv["course_id"], "course_name": inv["course_name"]}


# ---------- 全局管理（仅超级管理员） ----------

@app.get("/api/admin/users")
async def admin_list_users(uid: int = Depends(require_super_admin)):
    with db.get_db() as conn:
        rows = conn.execute(
            "SELECT id, username, display_name, role, created_at FROM users ORDER BY id"
        ).fetchall()
    return {"users": [dict(r) for r in rows]}


@app.delete("/api/admin/users/{user_id}")
async def admin_delete_user(user_id: int, uid: int = Depends(require_super_admin)):
    if user_id == uid:
        raise HTTPException(status_code=400, detail="不能删除自己")
    with db.get_db() as conn:
        cur = conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="用户不存在")
    return {"ok": True}


@app.get("/api/admin/courses")
async def admin_list_courses(uid: int = Depends(require_super_admin)):
    with db.get_db() as conn:
        rows = conn.execute(
            "SELECT c.id, c.name, u.username AS owner_username, c.created_at, "
            "(SELECT COUNT(*) FROM course_members m WHERE m.course_id = c.id) AS member_count "
            "FROM courses c JOIN users u ON u.id = c.owner_id ORDER BY c.id"
        ).fetchall()
    return {"courses": [dict(r) for r in rows]}


@app.delete("/api/admin/courses/{course_id}")
async def admin_delete_course(course_id: int, uid: int = Depends(require_super_admin)):
    with db.get_db() as conn:
        cur = conn.execute("DELETE FROM courses WHERE id = ?", (course_id,))
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="课程不存在")
    return {"ok": True}


# ---------- 班级 ----------

@app.get("/api/courses/{course_id}/classes")
async def list_classes(course_id: int, role: str = Depends(require_course_member)):
    with db.get_db() as conn:
        rows = conn.execute(
            "SELECT id, name, "
            "(SELECT COUNT(*) FROM students s WHERE s.class_id = classes.id) AS count "
            "FROM classes WHERE course_id = ? ORDER BY name",
            (course_id,),
        ).fetchall()
    return {"classes": [dict(r) for r in rows]}


@app.post("/api/courses/{course_id}/classes")
async def add_class(payload: dict, course_id: int, role: str = Depends(require_course_admin)):
    name = (payload.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="班级名不能为空")
    with db.get_db() as conn:
        try:
            cur = conn.execute(
                "INSERT INTO classes(name, course_id) VALUES (?, ?)", (name, course_id)
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=400, detail="该班级已存在")
    return {"id": cur.lastrowid, "name": name}


@app.delete("/api/courses/{course_id}/classes/{class_id}")
async def delete_class(course_id: int, class_id: int, role: str = Depends(require_course_admin)):
    with db.get_db() as conn:
        cur = conn.execute(
            "DELETE FROM classes WHERE id = ? AND course_id = ?", (class_id, course_id)
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="班级不存在")
    return {"ok": True}


# ---------- 实验 ----------

@app.get("/api/courses/{course_id}/experiments")
async def list_experiments(course_id: int, role: str = Depends(require_course_member)):
    with db.get_db() as conn:
        rows = conn.execute(
            "SELECT id, name FROM experiments WHERE course_id = ? ORDER BY id", (course_id,)
        ).fetchall()
    return {"experiments": [dict(r) for r in rows]}


@app.post("/api/courses/{course_id}/experiments")
async def add_experiment(payload: dict, course_id: int, role: str = Depends(require_course_admin)):
    name = (payload.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="实验名称不能为空")
    with db.get_db() as conn:
        try:
            cur = conn.execute(
                "INSERT INTO experiments(name, course_id) VALUES (?, ?)", (name, course_id)
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=400, detail="该实验已存在")
    return {"id": cur.lastrowid, "name": name}


@app.delete("/api/courses/{course_id}/experiments/{exp_id}")
async def delete_experiment(course_id: int, exp_id: int, role: str = Depends(require_course_admin)):
    with db.get_db() as conn:
        cur = conn.execute(
            "DELETE FROM experiments WHERE id = ? AND course_id = ?", (exp_id, course_id)
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="实验不存在")
    return {"ok": True}


# ---------- 导入 ----------

@app.post("/api/courses/{course_id}/import")
async def import_students(
    course_id: int,
    file: UploadFile = File(...),
    role: str = Depends(require_course_admin),
    class_id: int = 0,
):
    """导入学生名单。传 class_id 时全部归入该班级（忽略文件中的班级列）；
    不传时按文件中的班级列自动建班归类。"""
    data = await file.read()
    try:
        rows = importers.parse_file(file.filename, data)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    added_students = added_classes = skipped = conflicts = 0
    with db.get_db() as conn:
        if class_id:
            if not conn.execute(
                "SELECT id FROM classes WHERE id = ? AND course_id = ?", (class_id, course_id)
            ).fetchone():
                raise HTTPException(status_code=404, detail="班级不存在")
        for class_name, student_no, name in rows:
            if not student_no or not name:
                skipped += 1
                continue
            if class_id:
                cid = class_id
            else:
                cur = conn.execute(
                    "INSERT OR IGNORE INTO classes(name, course_id) VALUES (?, ?)",
                    (class_name, course_id),
                )
                if cur.rowcount > 0:
                    added_classes += 1
                cid = conn.execute(
                    "SELECT id FROM classes WHERE name = ? AND course_id = ?",
                    (class_name, course_id),
                ).fetchone()["id"]
            existing = conn.execute(
                "SELECT id, class_id FROM students WHERE student_no = ? AND course_id = ?",
                (student_no, course_id),
            ).fetchone()
            if existing is None:
                conn.execute(
                    "INSERT INTO students(class_id, student_no, name, course_id) "
                    "VALUES (?, ?, ?, ?)",
                    (cid, student_no, name, course_id),
                )
                added_students += 1
            elif existing["class_id"] == cid:
                conn.execute("UPDATE students SET name = ? WHERE id = ?", (name, existing["id"]))
                skipped += 1
            else:
                conflicts += 1
    return {
        "total": len(rows),
        "added_students": added_students,
        "added_classes": added_classes,
        "skipped": skipped,
        "conflicts": conflicts,
    }


# ---------- 学生 ----------

@app.get("/api/courses/{course_id}/students")
async def list_students(course_id: int, class_id: int, role: str = Depends(require_course_member)):
    with db.get_db() as conn:
        rows = conn.execute(
            "SELECT s.id, s.student_no, s.name FROM students s "
            "WHERE s.class_id = ? AND s.course_id = ? ORDER BY s.student_no",
            (class_id, course_id),
        ).fetchall()
    return {"students": [dict(r) for r in rows]}


@app.post("/api/courses/{course_id}/students")
async def add_student(payload: dict, course_id: int, role: str = Depends(require_course_admin)):
    class_id = payload.get("class_id")
    student_no = (payload.get("student_no") or "").strip()
    name = (payload.get("name") or "").strip()
    if not class_id or not student_no or not name:
        raise HTTPException(status_code=400, detail="班级、学号、姓名都不能为空")
    with db.get_db() as conn:
        if not conn.execute(
            "SELECT id FROM classes WHERE id = ? AND course_id = ?", (class_id, course_id)
        ).fetchone():
            raise HTTPException(status_code=404, detail="班级不存在")
        existing = conn.execute(
            "SELECT class_id FROM students WHERE student_no = ? AND course_id = ?",
            (student_no, course_id),
        ).fetchone()
        if existing:
            if existing["class_id"] == class_id:
                raise HTTPException(status_code=400, detail="该班级已存在相同学号的学生")
            raise HTTPException(status_code=400, detail="该学号已属于其他班级，不能重复")
        cur = conn.execute(
            "INSERT INTO students(class_id, student_no, name, course_id) VALUES (?, ?, ?, ?)",
            (class_id, student_no, name, course_id),
        )
    return {"id": cur.lastrowid, "student_no": student_no, "name": name}


@app.delete("/api/courses/{course_id}/students/{student_id}")
async def delete_student(course_id: int, student_id: int, role: str = Depends(require_course_admin)):
    with db.get_db() as conn:
        cur = conn.execute(
            "DELETE FROM students WHERE id = ? AND course_id = ?", (student_id, course_id)
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="学生不存在")
    return {"ok": True}


@app.get("/api/courses/{course_id}/student_lookup")
async def student_lookup(course_id: int, student_no: str, role: str = Depends(require_course_member)):
    student_no = (student_no or "").strip()
    if not student_no:
        raise HTTPException(status_code=400, detail="请输入学号")
    with db.get_db() as conn:
        stu = conn.execute(
            "SELECT s.id, s.student_no, s.name, c.name AS class_name "
            "FROM students s JOIN classes c ON c.id = s.class_id "
            "WHERE s.student_no = ? AND s.course_id = ?",
            (student_no, course_id),
        ).fetchone()
        if not stu:
            return {"found": False}
        experiments = conn.execute(
            "SELECT id, name FROM experiments WHERE course_id = ? ORDER BY id", (course_id,)
        ).fetchall()
        comps = conn.execute(
            "SELECT experiment_id, completed_at, operator FROM completions WHERE student_id = ?",
            (stu["id"],),
        ).fetchall()
    comp_map = {c["experiment_id"]: c for c in comps}
    exps_out = []
    for e in experiments:
        c = comp_map.get(e["id"])
        exps_out.append({
            "id": e["id"],
            "name": e["name"],
            "completed": c is not None,
            "completed_at": c["completed_at"] if c else None,
            "operator": c["operator"] if c else None,
        })
    return {
        "found": True,
        "student": {
            "id": stu["id"],
            "student_no": stu["student_no"],
            "name": stu["name"],
            "class_name": stu["class_name"],
        },
        "experiments": exps_out,
        "done_count": sum(1 for e in exps_out if e["completed"]),
    }


# ---------- 完成记录 ----------

@app.post("/api/courses/{course_id}/completions")
async def add_completion(
    payload: dict, request: Request, course_id: int, role: str = Depends(require_course_member)
):
    student_id = payload.get("student_id")
    experiment_id = payload.get("experiment_id")
    if not student_id or not experiment_id:
        raise HTTPException(status_code=400, detail="参数不完整")
    now = _now()
    operator = request.session.get("operator") or request.session.get("display_name", "")
    with db.get_db() as conn:
        if not conn.execute(
            "SELECT id FROM students WHERE id = ? AND course_id = ?", (student_id, course_id)
        ).fetchone():
            raise HTTPException(status_code=404, detail="学生不存在")
        if not conn.execute(
            "SELECT id FROM experiments WHERE id = ? AND course_id = ?", (experiment_id, course_id)
        ).fetchone():
            raise HTTPException(status_code=404, detail="实验不存在")
        try:
            conn.execute(
                "INSERT INTO completions(student_id, experiment_id, completed_at, operator) "
                "VALUES (?, ?, ?, ?)",
                (student_id, experiment_id, now, operator),
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=400, detail="该学生本次实验已登记，无需重复")
    return {"ok": True, "completed_at": now, "operator": operator}


@app.delete("/api/courses/{course_id}/completions")
async def delete_completion(
    student_id: int, experiment_id: int, course_id: int, role: str = Depends(require_course_member)
):
    with db.get_db() as conn:
        cur = conn.execute(
            "DELETE FROM completions WHERE student_id = ? AND experiment_id = ? "
            "AND student_id IN (SELECT id FROM students WHERE course_id = ?)",
            (student_id, experiment_id, course_id),
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="没有可撤销的记录")
    return {"ok": True}


@app.put("/api/courses/{course_id}/completions")
async def update_completion(
    payload: dict, request: Request, course_id: int, role: str = Depends(require_course_admin)
):
    student_id = payload.get("student_id")
    experiment_id = payload.get("experiment_id")
    completed_at = (payload.get("completed_at") or "").strip()
    if not student_id or not experiment_id or not completed_at:
        raise HTTPException(status_code=400, detail="参数不完整")
    try:
        dt = datetime.fromisoformat(completed_at)
        if dt.tzinfo is None:
            dt = dt.astimezone()
        completed_at = dt.isoformat(timespec="seconds")
    except ValueError:
        raise HTTPException(status_code=400, detail="时间格式不正确")
    operator = request.session.get("operator") or request.session.get("display_name", "")
    with db.get_db() as conn:
        cur = conn.execute(
            "UPDATE completions SET completed_at = ?, operator = ? "
            "WHERE student_id = ? AND experiment_id = ? "
            "AND student_id IN (SELECT id FROM students WHERE course_id = ?)",
            (completed_at, operator, student_id, experiment_id, course_id),
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="该学生此实验尚未登记，无法修改时间")
    return {"ok": True}


# ---------- 总览与导出 ----------

def _overview_data(class_id, course_id):
    """返回 (class_name, students, experiments, cell_map)。
    cell_map: {(student_id, experiment_id): {"completed_at": str, "operator": str}}"""
    with db.get_db() as conn:
        cls = conn.execute(
            "SELECT name FROM classes WHERE id = ? AND course_id = ?", (class_id, course_id)
        ).fetchone()
        students = conn.execute(
            "SELECT id, student_no, name FROM students "
            "WHERE class_id = ? AND course_id = ? ORDER BY student_no",
            (class_id, course_id),
        ).fetchall()
        experiments = conn.execute(
            "SELECT id, name FROM experiments WHERE course_id = ? ORDER BY id", (course_id,)
        ).fetchall()
        comps = conn.execute(
            "SELECT c.student_id, c.experiment_id, c.completed_at, c.operator "
            "FROM completions c JOIN students s ON s.id = c.student_id "
            "WHERE s.class_id = ? AND s.course_id = ?",
            (class_id, course_id),
        ).fetchall()
    cell_map = {}
    for c in comps:
        cell_map[(c["student_id"], c["experiment_id"])] = {
            "completed_at": c["completed_at"],
            "operator": c["operator"],
        }
    return (cls["name"] if cls else ""), students, experiments, cell_map


@app.get("/api/courses/{course_id}/overview")
async def overview(course_id: int, class_id: int, role: str = Depends(require_course_admin)):
    class_name, students, experiments, cell_map = _overview_data(class_id, course_id)
    if not class_name:
        raise HTTPException(status_code=404, detail="班级不存在")
    students_out = []
    for s in students:
        cells = {}
        done_count = 0
        for e in experiments:
            cell = cell_map.get((s["id"], e["id"]))
            cells[str(e["id"])] = cell
            if cell:
                done_count += 1
        students_out.append({
            "id": s["id"],
            "student_no": s["student_no"],
            "name": s["name"],
            "cells": cells,
            "done_count": done_count,
        })
    return {
        "class_name": class_name,
        "experiments": [dict(e) for e in experiments],
        "students": students_out,
    }


def _fmt_dt(iso):
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return iso


@app.get("/api/courses/{course_id}/export")
async def export_excel(course_id: int, class_id: int, role: str = Depends(require_course_admin)):
    class_name, students, experiments, cell_map = _overview_data(class_id, course_id)
    if not class_name:
        raise HTTPException(status_code=404, detail="班级不存在")

    wb = Workbook()
    ws = wb.active
    ws.title = "完成情况"

    header = ["学号", "姓名"] + [e["name"] for e in experiments] + ["已完成数"]
    ws.append(header)

    header_fill = PatternFill("solid", fgColor="2F7BFF")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    pending_fill = PatternFill("solid", fgColor="FDE8E8")
    done_font = Font(color="16A34A", bold=True)

    for s in students:
        row = [s["student_no"], s["name"]]
        done = 0
        for e in experiments:
            c = cell_map.get((s["id"], e["id"]))
            if c:
                row.append(_fmt_dt(c["completed_at"]))
                done += 1
            else:
                row.append("未完成")
        row.append(done)
        ws.append(row)

    last_col = 3 + len(experiments)
    for r in range(2, len(students) + 2):
        for col in range(3, 3 + len(experiments)):
            cell = ws.cell(row=r, column=col)
            if cell.value == "未完成":
                cell.fill = pending_fill
            else:
                cell.font = done_font
        ws.cell(row=r, column=last_col).alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    for i in range(len(experiments)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{class_name}_完成情况.xlsx" if class_name else "完成情况.xlsx"
    disposition = f"attachment; filename*=UTF-8''{quote(filename)}"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": disposition},
    )


# ---------- 静态文件 ----------

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/")
async def index():
    return FileResponse(os.path.join(STATIC_DIR, "index.html"))
